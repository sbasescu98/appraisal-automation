"""
Phase 2a: Universal Document Extraction & Chunking
Extracts text from downloaded planning documents (PDF, DOCX, RTF, TXT, optionally DOC)
and splits it into token-sized chunks for LLM processing.
"""

import asyncio
import json
import logging
import math
import re
import shutil
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from document_detector import DEFAULT_SUPPORTED_TYPES, detect_file_type, is_supported_type

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# pdfminer logs a warning for every malformed object it tolerates; keep the output readable
logging.getLogger('pdfminer').setLevel(logging.ERROR)

# Storage paths (match the scraper's layout)
DATA_DIR = Path("data")
DOCUMENTS_DIR = DATA_DIR / "documents"
METADATA_DIR = DATA_DIR / "metadata"
EXTRACTED_DIR = DATA_DIR / "extracted"

MAX_CHUNK_TOKENS = 4000

# Extra fields the search crawler (google_search_crawler.py) records for each result,
# carried through to raw_texts/*.json and all_documents.json so results can be retraced
SEARCH_FIELDS = (
    'source_url', 'final_url', 'title', 'snippet', 'displayed_link', 'result_date',
    'query', 'queries', 'search_rank', 'provider', 'search_id',
    'searched_at', 'scraped_at', 'content_type', 'mentions_project',
)

# Extra fields the Arcus scraper (arcus_scraper.py) records for each comment and document
COMMENT_FIELDS = ('commenter', 'stance', 'comment_id', 'created_at', 'category', 'file_id', 'uploaded_at')

EXTRA_SOURCE_FIELDS = SEARCH_FIELDS + COMMENT_FIELDS

# A PDF page with fewer extractable characters than this is treated as a scanned image
MIN_TEXT_CHARS_PER_PAGE = 25

# Fonts without a unicode map come out as "(cid:36)(cid:81)..." placeholders. A page where
# more than this share of the text is placeholders is unreadable and gets OCR'd instead.
MAX_CID_RATIO = 0.2
CID_PATTERN = re.compile(r'\(cid:\d+\)')

# A web page with less main text than this is a loading screen or an empty shell, not content
MIN_HTML_TEXT_CHARS = 150

# Resolution used when rendering scanned pages for OCR
OCR_DPI = 300


# ---------------------------------------------------------------------------
# Tokens & chunking
# ---------------------------------------------------------------------------

def estimate_tokens(text: str) -> int:
    """Rough token estimate (~4 characters per token for English text)."""
    return math.ceil(len(text) / 4) if text else 0


def chunk_by_tokens(text: str, max_tokens: int = MAX_CHUNK_TOKENS) -> list[dict]:
    """
    Split text into chunks of at most ~max_tokens.

    Splits on paragraph boundaries where possible, then on sentences,
    and finally on words for anything still too long.

    Returns:
        List of {'chunk_id', 'text', 'estimated_tokens'} dicts
    """
    if not text or not text.strip():
        return []

    max_chars = max_tokens * 4

    # Break the text into pieces that each fit within one chunk
    pieces = []
    for para in re.split(r'\n\s*\n', text):
        para = para.strip()
        if not para:
            continue
        if len(para) <= max_chars:
            pieces.append(para)
            continue
        for sentence in re.split(r'(?<=[.!?])\s+', para):
            if len(sentence) <= max_chars:
                pieces.append(sentence)
                continue
            # A single huge "sentence" (e.g. a table dump): split on words
            current = ''
            for word in sentence.split():
                if current and len(current) + 1 + len(word) > max_chars:
                    pieces.append(current)
                    current = word
                else:
                    current = f"{current} {word}" if current else word
            if current:
                pieces.append(current)

    # Pack pieces greedily into chunks
    chunks = []
    current = ''
    for piece in pieces:
        if current and len(current) + 2 + len(piece) > max_chars:
            chunks.append(current)
            current = piece
        else:
            current = f"{current}\n\n{piece}" if current else piece
    if current:
        chunks.append(current)

    return [
        {'chunk_id': i, 'text': chunk, 'estimated_tokens': estimate_tokens(chunk)}
        for i, chunk in enumerate(chunks)
    ]


def _error_result(file_path: Path, file_type: str, error: str, status: str = 'error') -> dict:
    return {
        'filename': file_path.name,
        'file_type': file_type,
        'status': status,
        'error': error,
        'text': None,
    }


def _success_result(file_path: Path, file_type: str, text: str, metadata: dict, **extra) -> dict:
    return {
        'filename': file_path.name,
        'file_type': file_type,
        **extra,
        'status': 'success',
        'text': text,
        'chunks': chunk_by_tokens(text),
        'metadata': {**metadata, 'estimated_tokens': estimate_tokens(text)},
    }


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def ocr_available() -> bool:
    """OCR needs the tesseract binary and poppler's pdftoppm (used by pdf2image)."""
    return bool(shutil.which('tesseract') and shutil.which('pdftoppm'))


def count_pdf_pages(file_path: Path) -> int:
    import pdfplumber
    with pdfplumber.open(file_path) as pdf:
        return len(pdf.pages)


def _extract_pdf_pages(file_path: Path) -> list[str]:
    """Extract the embedded text of each page (empty string for image-only pages)."""
    import pdfplumber
    with pdfplumber.open(file_path) as pdf:
        return [(page.extract_text() or '').strip() for page in pdf.pages]


def _page_needs_ocr(text: str) -> bool:
    """True if a page has no usable text layer (image-only, or unmapped font placeholders)."""
    cid_chars = sum(len(m) for m in CID_PATTERN.findall(text))
    if text and cid_chars / len(text) > MAX_CID_RATIO:
        return True
    return len(CID_PATTERN.sub('', text).strip()) < MIN_TEXT_CHARS_PER_PAGE


def _classify_pages(page_texts: list[str]) -> str:
    """Return 'text', 'scanned' or 'mixed' depending on how many pages lack usable text."""
    if not page_texts:
        return 'empty'
    text_pages = sum(not _page_needs_ocr(t) for t in page_texts)
    if text_pages == len(page_texts):
        return 'text'
    if text_pages == 0:
        return 'scanned'
    return 'mixed'


async def analyze_pdf(file_path: Path) -> str:
    """
    Detect whether a PDF has a text layer.

    Returns:
        'text' (all pages have text), 'scanned' (no pages do),
        'mixed' (some do) or 'empty' (no pages)
    """
    page_texts = await asyncio.to_thread(_extract_pdf_pages, file_path)
    return _classify_pages(page_texts)


def _ocr_pdf_pages(file_path: Path, page_numbers: list[int]) -> dict[int, str]:
    """OCR the given 1-based page numbers, one page at a time to limit memory use."""
    import pytesseract
    from pdf2image import convert_from_path

    results = {}
    for page_num in page_numbers:
        images = convert_from_path(str(file_path), dpi=OCR_DPI, first_page=page_num, last_page=page_num)
        results[page_num] = pytesseract.image_to_string(images[0]).strip() if images else ''
    return results


async def extract_pdf(file_path: Path) -> dict:
    """
    Extract text from a PDF.

    Pages with a text layer are read with pdfplumber; pages without one
    (scanned images) are OCR'd with tesseract when it is installed.
    """
    try:
        page_texts = await asyncio.to_thread(_extract_pdf_pages, file_path)
    except Exception as e:
        return _error_result(file_path, 'pdf', f"Could not read PDF: {e}")

    pdf_subtype = _classify_pages(page_texts)
    if pdf_subtype == 'empty':
        return _error_result(file_path, 'pdf', 'PDF has no pages', status='skipped')

    # Pages that need OCR (1-based)
    ocr_pages = [i + 1 for i, t in enumerate(page_texts) if _page_needs_ocr(t)]
    ocr_skipped = []

    if ocr_pages:
        if ocr_available():
            try:
                ocr_texts = await asyncio.to_thread(_ocr_pdf_pages, file_path, ocr_pages)
                for page_num, text in ocr_texts.items():
                    page_texts[page_num - 1] = text
            except Exception as e:
                if pdf_subtype == 'scanned':
                    return _error_result(file_path, 'pdf', f"OCR failed: {e}")
                logger.warning(f"OCR failed for {file_path.name}, keeping text pages only: {e}")
                ocr_skipped = ocr_pages
        else:
            if pdf_subtype == 'scanned':
                return _error_result(
                    file_path, 'pdf',
                    'Scanned PDF needs OCR, but tesseract/poppler are not installed '
                    '(brew install tesseract poppler)'
                )
            logger.warning(
                f"{file_path.name}: {len(ocr_pages)} image-only pages skipped "
                f"(install tesseract + poppler for OCR)"
            )
            ocr_skipped = ocr_pages

    # Drop pages we couldn't OCR, and any stray font placeholders left on text pages
    page_texts = [
        '' if (i + 1) in ocr_skipped else CID_PATTERN.sub('', t).strip()
        for i, t in enumerate(page_texts)
    ]
    
    # Keep page boundaries visible so later phases can cite pages
    text = '\n\n'.join(
        f"[Page {i}]\n{t}" for i, t in enumerate(page_texts, 1) if t
    )

    ocr_used = bool(ocr_pages) and not ocr_skipped
    return _success_result(
        file_path, 'pdf', text,
        {
            'extraction_method': 'pdfplumber+pytesseract' if ocr_used and pdf_subtype == 'mixed'
                                 else 'pytesseract' if ocr_used else 'pdfplumber',
            'total_pages': len(page_texts),
            'ocr_used': ocr_used,
            'ocr_pages': ocr_pages if ocr_used else [],
            'ocr_pages_skipped': ocr_skipped,
        },
        pdf_subtype=pdf_subtype,
    )


# ---------------------------------------------------------------------------
# DOCX / DOC
# ---------------------------------------------------------------------------

def _extract_docx_sync(file_path: Path) -> tuple[str, int, int]:
    import docx

    doc = docx.Document(str(file_path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # Tables as pipe-separated rows
    for table in doc.tables:
        for row in table.rows:
            parts.append(' | '.join(cell.text.strip() for cell in row.cells))

    return '\n'.join(parts), len(doc.paragraphs), len(doc.tables)


async def extract_docx(file_path: Path) -> dict:
    """Extract text (paragraphs and tables) from a DOCX file."""
    try:
        text, num_paragraphs, num_tables = await asyncio.to_thread(_extract_docx_sync, file_path)
    except Exception as e:
        return _error_result(file_path, 'docx', str(e))

    return _success_result(file_path, 'docx', text, {
        'extraction_method': 'python-docx',
        'num_paragraphs': num_paragraphs,
        'num_tables': num_tables,
    })


async def extract_doc(file_path: Path) -> dict:
    """
    Extract text from a legacy .doc file using macOS's built-in textutil.
    (Not enabled by default; pass 'doc' in supported_types.)
    """
    if not shutil.which('textutil'):
        return _error_result(file_path, 'doc', '.doc extraction requires macOS textutil')

    try:
        proc = await asyncio.to_thread(
            subprocess.run,
            ['textutil', '-convert', 'txt', '-stdout', str(file_path)],
            capture_output=True, text=True, timeout=120,
        )
        if proc.returncode != 0:
            raise RuntimeError(proc.stderr.strip() or f"textutil exited with {proc.returncode}")
    except Exception as e:
        return _error_result(file_path, 'doc', str(e))

    return _success_result(file_path, 'doc', proc.stdout.strip(), {'extraction_method': 'textutil'})


# ---------------------------------------------------------------------------
# RTF / TXT
# ---------------------------------------------------------------------------

async def extract_rtf(file_path: Path) -> dict:
    """Extract text from an RTF file by stripping its formatting codes."""
    from striprtf.striprtf import rtf_to_text

    try:
        rtf_content = file_path.read_text(encoding='utf-8', errors='ignore')
        plain_text = rtf_to_text(rtf_content, errors='ignore')
    except Exception as e:
        return _error_result(file_path, 'rtf', str(e))

    lines = [line.strip() for line in plain_text.split('\n') if line.strip()]
    return _success_result(file_path, 'rtf', '\n'.join(lines), {'extraction_method': 'striprtf'})


async def extract_txt(file_path: Path) -> dict:
    """Read a plain text file, trying common encodings."""
    try:
        raw = file_path.read_bytes()
    except Exception as e:
        return _error_result(file_path, 'txt', str(e))

    # latin-1 accepts any byte sequence, so this always finds an encoding
    for encoding in ['utf-8', 'cp1252', 'latin-1']:
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    return _success_result(file_path, 'txt', text, {
        'extraction_method': 'raw_read',
        'encoding': encoding,
    })


# ---------------------------------------------------------------------------
# HTML (web pages from the search crawler)
# ---------------------------------------------------------------------------

def _extract_html_sync(file_path: Path) -> tuple[Optional[str], dict]:
    import trafilatura

    raw = file_path.read_bytes()
    # Main content only: navigation, footers, cookie banners and comments are dropped
    text = trafilatura.extract(raw, include_tables=True, include_comments=False, favor_recall=True)
    meta = trafilatura.extract_metadata(raw)
    page_meta = {
        'page_title': meta.title if meta else None,
        'page_author': meta.author if meta else None,
        'page_date': meta.date if meta else None,
        'site_name': meta.sitename if meta else None,
    }
    return text, page_meta


async def extract_html(file_path: Path) -> dict:
    """Extract the main text of a web page with trafilatura."""
    try:
        text, page_meta = await asyncio.to_thread(_extract_html_sync, file_path)
    except Exception as e:
        return _error_result(file_path, 'html', str(e))

    if not text or len(text.strip()) < MIN_HTML_TEXT_CHARS:
        return _error_result(file_path, 'html', 'No main text found on page')

    return _success_result(file_path, 'html', text.strip(), {'extraction_method': 'trafilatura', **page_meta})


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------

EXTRACTORS = {
    'pdf': extract_pdf,
    'docx': extract_docx,
    'doc': extract_doc,
    'rtf': extract_rtf,
    'txt': extract_txt,
    'html': extract_html,
}


async def extract_single_document(file_path: Path) -> dict:
    """Unified entry point for any document type; routes to the right extractor."""
    doc_type = detect_file_type(file_path)
    extractor = EXTRACTORS.get(doc_type)

    if extractor is None:
        return _error_result(file_path, doc_type, f'Unsupported file type: {doc_type}', status='skipped')

    return await extractor(file_path)


# ---------------------------------------------------------------------------
# Extractor class
# ---------------------------------------------------------------------------

class DocumentExtractor:
    """
    Universal document extractor.
    Handles PDF, DOCX, RTF, TXT, and optionally DOC.
    """

    def __init__(self, application_key: str, supported_types: Optional[list] = None):
        self.application_key = application_key
        self.docs_dir = DOCUMENTS_DIR / application_key
        self.extract_dir = EXTRACTED_DIR / application_key
        self.raw_texts_dir = self.extract_dir / "raw_texts"
        self.raw_texts_dir.mkdir(parents=True, exist_ok=True)

        # Which file types to extract ('doc' is opt-in)
        self.supported_types = supported_types or list(DEFAULT_SUPPORTED_TYPES)

        self.results = {}
        self._source_info = self._load_source_info()

    def _load_source_info(self) -> dict:
        """Map downloaded filenames to the scraper's table data (type, description, date, URL)."""
        metadata_file = METADATA_DIR / f"{self.application_key}_metadata.json"
        if not metadata_file.exists():
            return {}

        with open(metadata_file) as f:
            metadata = json.load(f)

        table = {d['filename']: d for d in metadata.get('documents_table', [])}
        info = {}
        for d in metadata.get('downloaded', []):
            row = table.get(d['filename'], {})
            merged = {**row, **d}
            info[Path(d['file_path']).name] = {
                'document_type': row.get('document_type'),
                'description': row.get('description'),
                'date_published': row.get('date_published'),
                'drawing_number': row.get('drawing_number'),
                'view_url': row.get('view_url'),
                **{f: merged[f] for f in EXTRA_SOURCE_FIELDS if f in merged},
            }
        return info

    def list_documents(self) -> list[Path]:
        if not self.docs_dir.exists():
            raise FileNotFoundError(
                f"No documents folder at {self.docs_dir}. Run the scraper first "
                f"(python cli.py -k KEY or --url URL); it prints the key to use, e.g. eastriding_25_01104_plf"
            )
        return sorted(p for p in self.docs_dir.iterdir() if p.is_file() and not p.name.startswith('.'))

    async def extract_all_documents(self) -> dict:
        """
        Extract every supported document in docs_dir, save one JSON per
        document plus an extraction_summary.json, and return the summary.
        """
        files = []

        documents = self.list_documents()

        # Output is named after the file stem; include the extension when two files share a stem
        stems = [p.stem for p in documents]
        self._clashing_stems = {stem for stem in stems if stems.count(stem) > 1}

        for file_path in documents:
            doc_type = detect_file_type(file_path)

            if not is_supported_type(doc_type, self.supported_types):
                logger.warning(f"Skipping unsupported: {file_path.name} ({doc_type})")
                result = _error_result(file_path, doc_type, f'Unsupported file type: {doc_type}', status='skipped')
            else:
                logger.info(f"Extracting {file_path.name}")
                result = await extract_single_document(file_path)

            result['source'] = self._source_info.get(file_path.name)
            self.results[file_path.name] = result

            entry = {
                'filename': result['filename'],
                'file_type': result['file_type'],
                'status': result['status'],
            }
            if result['status'] == 'success':
                output_file = self._save_extraction(result)
                entry['output_file'] = str(output_file.relative_to(self.extract_dir))
                entry['extracted_tokens'] = result['metadata']['estimated_tokens']
                entry['num_chunks'] = len(result['chunks'])
                if 'pdf_subtype' in result:
                    entry['pdf_subtype'] = result['pdf_subtype']
                logger.info(
                    f"  ✓ {result['metadata']['estimated_tokens']:,} tokens, {len(result['chunks'])} chunks"
                )
            else:
                entry['error'] = result['error']
                if result['status'] == 'error':
                    logger.error(f"  ✗ {result['error']}")
            files.append(entry)

        summary = self._generate_summary(files)
        self._save_summary(summary)
        self._log_extraction_summary(summary)
        self.compile_extractions()
        return summary

    def compile_extractions(self) -> tuple[Path, Path]:
        """
        Combine every successful extraction into all_documents.json and a
        readable all_documents.txt, in the order listed in extraction_summary.json.
        Works from the saved files, so it can be re-run without re-extracting.
        """
        summary_file = self.extract_dir / "extraction_summary.json"
        if not summary_file.exists():
            raise FileNotFoundError(
                f"No extraction summary at {summary_file}. Run: python extractor_cli.py extract {self.application_key}"
            )
        with open(summary_file, encoding='utf-8') as f:
            summary = json.load(f)

        documents = []
        for entry in summary['files']:
            if entry['status'] != 'success':
                continue
            with open(self.extract_dir / entry['output_file'], encoding='utf-8') as f:
                result = json.load(f)
            source = result.get('source') or {}
            documents.append({
                'filename': result['filename'],
                'file_type': result['file_type'],
                'document_type': source.get('document_type'),
                'description': source.get('description'),
                'date_published': source.get('date_published'),
                'view_url': source.get('view_url'),
                'total_pages': result['metadata'].get('total_pages'),
                'estimated_tokens': result['metadata']['estimated_tokens'],
                **{f: source[f] for f in EXTRA_SOURCE_FIELDS if f in source},
                'text': result['text'],
            })

        json_file, txt_file = self._write_compiled('all_documents', documents)

        # Applications with public comments (Arcus registers) also get them split out
        comments = [d for d in documents if d['document_type'] == 'Comment']
        if comments:
            self._write_compiled('comments', comments)
            self._write_compiled('documents', [d for d in documents if d['document_type'] != 'Comment'])

        logger.info(f"Compiled {len(documents)} documents into {json_file} and {txt_file}")
        return json_file, txt_file

    def _write_compiled(self, name: str, documents: list[dict]) -> tuple[Path, Path]:
        """Write {name}.json and a readable {name}.txt for a list of compiled documents."""
        is_comments = name == 'comments'

        json_file = self.extract_dir / f"{name}.json"
        with open(json_file, 'w', encoding='utf-8') as f:
            json.dump({
                'application_key': self.application_key,
                f'total_{name if is_comments else "documents"}': len(documents),
                'total_tokens': sum(d['estimated_tokens'] for d in documents),
                name if is_comments else 'documents': documents,
            }, f, indent=2, ensure_ascii=False)

        txt_file = self.extract_dir / f"{name}.txt"
        with open(txt_file, 'w', encoding='utf-8') as f:
            if is_comments:
                stances = {}
                for doc in documents:
                    stances[doc.get('stance') or 'Unknown'] = stances.get(doc.get('stance') or 'Unknown', 0) + 1
                f.write(f"APPLICATION {self.application_key} - {len(documents)} comments\n")
                f.write(', '.join(f"{n} {s}" for s, n in sorted(stances.items(), key=lambda x: -x[1])) + "\n")
                if documents and documents[0].get('view_url'):
                    f.write(f"Source: {documents[0]['view_url']}\n")
                # Each comment's text starts with its own header (commenter, stance, date)
                for i, doc in enumerate(documents, 1):
                    f.write(f"\n\n{'-' * 30} COMMENT {i}/{len(documents)} {'-' * 30}\n")
                    f.write(doc['text'].strip())
                    f.write("\n")
                return json_file, txt_file

            f.write(f"APPLICATION {self.application_key} - {len(documents)} documents\n")
            for i, doc in enumerate(documents, 1):
                f.write(f"\n\n{'=' * 80}\n")
                f.write(f"DOCUMENT {i}/{len(documents)}: {doc['filename']}\n")
                if doc.get('source_url'):
                    f.write(f"Title: {doc['title']}\n")
                    f.write(f"Query: {doc['query']} | Rank: {doc['search_rank']} | "
                            f"Retrieved: {doc.get('scraped_at') or ''}\n")
                elif doc['document_type']:
                    f.write(f"Type: {doc['document_type']} | Description: {doc['description']} | "
                            f"Published: {doc['date_published']}\n")
                if doc['view_url']:
                    f.write(f"Source: {doc['view_url']}\n")
                f.write(f"{'=' * 80}\n\n")
                f.write(doc['text'])
                f.write("\n")
        return json_file, txt_file

    def _save_extraction(self, result: dict) -> Path:
        path = Path(result['filename'])
        name = f"{path.stem}_{path.suffix.lstrip('.')}" if path.stem in self._clashing_stems else path.stem
        output_file = self.raw_texts_dir / f"{name}.json"
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(result, f, indent=2, ensure_ascii=False)
        return output_file

    def _generate_summary(self, files: list[dict]) -> dict:
        by_type = {}
        for entry in files:
            if entry['status'] == 'skipped' and entry['file_type'] not in self.supported_types:
                continue
            stats = by_type.setdefault(
                entry['file_type'], {'count': 0, 'successful': 0, 'failed': 0, 'total_tokens': 0}
            )
            stats['count'] += 1
            if entry['status'] == 'success':
                stats['successful'] += 1
                stats['total_tokens'] += entry['extracted_tokens']
            else:
                stats['failed'] += 1

        successful = [r for r in self.results.values() if r['status'] == 'success']
        return {
            'application_key': self.application_key,
            'extraction_date': datetime.now(timezone.utc).isoformat(),
            'total_documents': len(files),
            'by_type': by_type,
            'overall': {
                'successful': len(successful),
                'failed': sum(e['status'] == 'error' for e in files),
                'skipped': sum(e['status'] == 'skipped' for e in files),
                'total_chars': sum(len(r['text']) for r in successful),
                'total_tokens': sum(r['metadata']['estimated_tokens'] for r in successful),
            },
            'files': files,
        }

    def _save_summary(self, summary: dict) -> None:
        summary_file = self.extract_dir / "extraction_summary.json"
        with open(summary_file, 'w', encoding='utf-8') as f:
            json.dump(summary, f, indent=2, ensure_ascii=False)
        logger.info(f"Saved summary to {summary_file}")

    def _log_extraction_summary(self, summary: dict) -> None:
        overall = summary['overall']
        logger.info(
            f"Extracted {overall['successful']}/{summary['total_documents']} documents "
            f"({overall['failed']} failed, {overall['skipped']} skipped), "
            f"{overall['total_tokens']:,} tokens"
        )


async def main():
    """Example usage."""
    extractor = DocumentExtractor("P7SRNYBJ0GJ00")
    await extractor.extract_all_documents()


if __name__ == "__main__":
    asyncio.run(main())
